from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook


def buat_laporan(path, data):
    buku = Workbook()
    lembar = buku.active
    lembar.append(["Nama", "Nilai"])
    for nama, nilai in data:
        lembar.append([nama, nilai])
    lembar.append(["Rata-rata", f"=AVERAGE(B2:B{len(data) + 1})"])
    buku.save(path)


if __name__ == "__main__":
    # openpyxl menulis rumus tetapi tidak menghitung hasilnya (butuh aplikasi spreadsheet).
    with TemporaryDirectory() as folder:
        path = Path(folder) / "laporan.xlsx"
        buat_laporan(path, [("Ayu", 90), ("Budi", 80)])
        buku = load_workbook(path, read_only=True)
        print("Baris:", list(buku.active.values))
        buku.close()