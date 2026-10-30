from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

with TemporaryDirectory() as folder:
    path = Path(folder) / "nilai.xlsx"
    buku = Workbook()
    buku.active.append(["Nama", "Nilai"])
    buku.active.append(["Ayu", 90])
    buku.save(path)
    # Muat XLSX sungguhan untuk memverifikasi data setelah disimpan.
    hasil = load_workbook(path, read_only=True)
    print("Baris:", list(hasil.active.values))
    hasil.close()